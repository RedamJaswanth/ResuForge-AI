import io
import re
from typing import Tuple, Optional
import pypdf

try:
    import pdfplumber
    HAS_PDFPLUMBER = True
except ImportError:
    HAS_PDFPLUMBER = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

def clean_extracted_text(text: str) -> str:
    """Normalize extracted text whitespace and characters."""
    if not text:
        return ""
    # Normalize excessive newlines and spaces
    text = re.sub(r'\r\n|\r', '\n', text)
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF using pdfplumber if available, with pypdf fallback."""
    text_content = []
    
    if HAS_PDFPLUMBER:
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text(layout=True) or page.extract_text()
                    if page_text:
                        text_content.append(page_text)
            full_text = "\n\n".join(text_content).strip()
            if full_text:
                return clean_extracted_text(full_text)
        except Exception:
            pass  # Fallback to pypdf

    # Fallback to pypdf
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text_content.append(t)
        return clean_extracted_text("\n\n".join(text_content))
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX file."""
    if not HAS_DOCX:
        raise ImportError("python-docx is not installed")
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return clean_extracted_text("\n".join(paragraphs))
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")

def parse_uploaded_file(filename: str, file_bytes: bytes) -> Tuple[str, str]:
    """Parse uploaded file and return (extracted_text, file_type)."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes), "pdf"
    elif lower.endswith(".docx") or lower.endswith(".doc"):
        return extract_text_from_docx(file_bytes), "docx"
    elif lower.endswith(".txt"):
        return clean_extracted_text(file_bytes.decode("utf-8", errors="ignore")), "txt"
    else:
        raise ValueError("Unsupported file format. Please upload a PDF, DOCX, or TXT file.")
